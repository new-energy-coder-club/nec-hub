# -*- coding: utf-8 -*-
"""生成 NEC-HUB 整站评审与量化优化清单 Word 文档"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

def set_ea(run, font='微软雅黑'):
    run.font.name = 'Calibri'
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), font)

# ---------- 全局字体 ----------
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(10.5)
style.element.get_or_add_rPr()
rfonts = style.element.rPr.find(qn('w:rFonts'))
if rfonts is None:
    rfonts = OxmlElement('w:rFonts')
    style.element.rPr.append(rfonts)
rfonts.set(qn('w:eastAsia'), '微软雅黑')

def set_cell_bg(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), color)
    tcPr.append(shd)

def add_heading(text, level):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        set_ea(run)
        run.font.color.rgb = RGBColor(0x1F, 0x3B, 0x63)
    return h

def add_para(text, bold=False, size=10.5):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    return p

# ---------- 封面 ----------
for _ in range(6):
    doc.add_paragraph()
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run('NEC-HUB 整站体验评审与量化优化清单')
r.bold = True
r.font.size = Pt(22)
set_ea(r)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run('（首页 / 硬件评审 / MES 库位 / 合同审批）')
r.font.size = Pt(14)
set_ea(r)

doc.add_paragraph()
info = doc.add_paragraph()
info.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = info.add_run('版本：V1.0    日期：2026-10-14    状态：硬件评审章完成，其余章节待合并补评')
r.font.size = Pt(10.5)
set_ea(r)
doc.add_page_break()

# ---------- 评审口径 ----------
add_heading('评审口径与量化说明', 1)
add_para('优先级定义：', bold=True)
add_para('P0 — 阻塞使用的硬伤，不修复则核心流程无法顺利完成；建议立即排期。')
add_para('P1 — 一致性/联动/导航层问题，显著影响效率与正确性，建议本轮迭代修复。')
add_para('P2 — 体验与权限打磨项，可排入后续迭代。')
add_para('D — 数据联动类需求，涉及跨模块数据回写，需单独评估依赖后再排期。')
doc.add_paragraph()
add_para('量化口径：', bold=True)
add_para('每项给出“预估工时（人日）”，按一名熟悉本项目技术栈的前端全栈工程师独立完成的估算，'
         '含自测不含联调排期等待；实际排期请结合人力与依赖项再压缩或放宽。')
add_para('每项给出“验收标准”，作为完成定义（DoD），测试/验收时逐条核对。')

# ---------- 第一章：硬件评审 ----------
doc.add_page_break()
add_heading('第一章  硬件评审页（本次评审）', 1)
add_para('评审对象：nec-hub「硬件评审」提交页（/apps/hardware-review）。'
         '页面骨架完整（表单字段齐、有参数模板与批量导入能力），核心问题集中在：流程不可见、'
         '表单无状态管理、进阶功能裸奔给所有用户。共 13 项，其中数据联动项 D-13 含 3 个子项。')

COLS = ['编号', '优先级', '问题', '影响', '优化建议', '预估工时(人日)', '验收标准']
WIDTHS = [Cm(1.4), Cm(1.4), Cm(4.2), Cm(3.6), Cm(5.2), Cm(1.8), Cm(4.6)]

items = [
    # 编号, 优先级, 问题, 影响, 建议, 工时, 验收
    ('P0-01', 'P0',
     '「硬件评审流程」卡片为空：卡片名叫"流程"，点进去直接就是表单，用户看不到评审有几步、谁审、要多久。',
     '提交人对评审周期与责任方零预期，易反复催问；流程透明度缺失。',
     '表单上方加流程步骤条：提交 → 形式审查 → 专业评审 → 结论（通过/有条件通过/驳回），每步标注责任角色与 SLA（如形式审查 1 个工作日）。',
     '1.5',
     '步骤条在表单页首屏可见；每步显示角色与 SLA；当前步骤可随单状态高亮。'),
    ('P0-02', 'P0',
     '「提交」按钮与表单脱节：按钮悬在空卡片顶部，长表单提交时不在视野内；且无草稿能力，误刷新全丢。',
     '长表单填写成本高，误操作导致全部丢失，直接阻塞使用。',
     '表单底部加固定操作栏：保存草稿 + 提交评审双按钮；草稿本地+服务端双写，支持断点续填。',
     '2.0',
     '滚动至任意位置均可在底部固定栏提交/存草稿；草稿可恢复、可多份；刷新/重进不丢数据。'),
    ('P0-03', 'P0',
     '提交后无反馈闭环：看不到评审单号、预计完成时间、查询入口。',
     '提交人无法跟踪进度，重复提交或线下追问。',
     '提交成功页三件套：单号（建议规则 HR-2026-001，与项目编号 PRJ- 区分）、状态查询链接、评审人信息。',
     '1.0',
     '提交成功即展示三件套；单号全局唯一可检索；状态页信息与评审端一致。'),
    ('P1-04', 'P1',
     '必填标记不统一：硬件类型、规格参数未标必填，而二者恰是评审路由关键（类型决定评审人分配，参数是评审主体）。',
     '缺类型/参数的评审单无法正确路由，评审人被分错或打回重填。',
     '硬件类型、规格参数改为必填（红星标记），提交前校验拦截。',
     '0.3',
     '两项缺失时提交被拦截并定位到缺失字段；与既有必填项视觉一致。'),
    ('P1-05', 'P1',
     '项目编号纯手填：仅 placeholder 示例（PRJ-2026-001），无格式校验、无查重、不关联已有项目。',
     '手填串号、虚构编号，评审与任务排期/赛季目标数据断裂。',
     '改为下拉选择已有项目，自动带出编号；同时保留格式校验与查重兜底。',
     '2.5',
     '只能选已有项目或按规则新建；编号自动带出不可手改；与任务排期数据可互相跳转。'),
    ('P1-06', 'P1',
     '硬件类型与参数模板不联动："按类型填充常用参数"是弱化文字链接，藏在表格右上角。',
     '参数填写随意，评审口径不一致，评审人来回补问。',
     '选完类型自动预填常用参数行，链接升级为按钮并保留"追加填充"；内置模板：主控板（主控芯片/工作电压/功耗/通信接口/工作温度）、激光雷达（量程/精度/视场角/点频/接口）、驱动电机（额定功率/电压/转速/扭矩/绝缘等级）。',
     '1.5',
     '切换类型即预填对应模板；已有手填行不被覆盖（追加模式）；模板可维护。'),
    ('P1-07', 'P1',
     '上传文件无后续管理：只有上传入口，无已传文件的预览、替换、删除、版本记录。',
     '驳回后改版（V1.0→V1.1）无法并存可查，评审追溯困难。',
     '附件区加版本管理：预览、替换、删除、版本历史（V1.0/V1.1 并存，标注当前生效版）。',
     '3.0',
     '同一资料可多版本并存；每版可下载预览；评审单上标注当前生效版本。'),
    ('P1-08', 'P1',
     '一级导航溢出：11 个导航项 + 灰色角标，"硬件评审"截断为"硬件评"，无横向滚动视觉提示。',
     '导航项不可达/不可读，信息架构失控。',
     '按角标逻辑收拢为三组下拉：工作台（我的首页/我的任务/报销/任务排期/赛季目标）、资源（应用广场/Skill中心/资料页）、评审（图纸评审/开源仓库/硬件评审）；"评审"角标保留，"资源"角标去掉。',
     '1.0',
     '三组下拉在 1280px 宽度下无截断；任意导航项两次点击内可达；角标与分组一致。'),
    ('P2-09', 'P2',
     '「批量导入」暴露技术术语：名称|文件token|链接格式面向开发者，普通提交人看到"token"会懵。',
     '普通用户误用/困惑；技术字段裸奔有安全风险。',
     '折叠为"高级导入（维护者）"，按角色显隐；普通用户仅保留拖拽上传。',
     '0.5',
     '非维护者角色看不到批量导入入口；维护者导入功能不受影响。'),
    ('P2-10', 'P2',
     '单文件 ≤10MB 对硬件格式偏小：STEP/点云易超；且无格式校验提示。',
     '大文件上传失败无引导，反复试错。',
     '放宽至 50MB 走分片上传，或上传区直接引导"超大文件传网盘贴链接"；补格式校验提示（DXF 版本、STEP AP203/214）。',
     '2.0',
     '50MB 内文件稳定上传；超限给出明确引导文案；不符合版本要求的文件在上传前被提示。'),
    ('P2-11', 'P2',
     '评审标准不公开：提交人不知道按什么清单被审。',
     '提交质量参差，评审一轮通过率低的隐性原因。',
     '公示《硬件评审检查表》（原理图完整性/电源设计/接口防护/EMC 预留/可制造性），挂在流程步骤条下方，提交前可自查。',
     '1.0',
     '检查表在表单页可见可展开；条目与评审人实际使用的清单一致。'),
    ('P2-12', 'P2',
     '「代他人提交」缺选人机制：点开之后如何选目标人、权限边界均不可见。',
     '代提交功能形同虚设，或越权代提。',
     '加人员搜索 + 说明文案（谁能代谁提：如组长可代组员、管理员可代全员）。',
     '1.5',
     '可搜索选择目标人；无权限时给出边界说明；被代提人可在本人列表中看到该单。'),
    ('D-13a', 'D',
     '数据联动①：项目编号 ↔ 任务排期/赛季目标数据未打通。',
     '评审与排期两套数据，项目状态需人工对齐。',
     '项目编号与任务排期联动，评审状态回写至排期视图。',
     '1.5',
     '排期页可查看关联评审单及状态；状态变更自动回写。'),
    ('D-13b', 'D',
     '数据联动②：硬件评审单 ↔ 图纸评审单未关联（同一块板的原理图与 PCB 各自过审）。',
     '同一硬件的多类评审单孤立，无法总览整体进度。',
     '评审单间建立关联（同一硬件聚合展示图纸评审 + 硬件评审状态）。',
     '1.0',
     '从硬件评审单可跳转/查看关联图纸评审单及其状态。'),
    ('D-13c', 'D',
     '数据联动③：评审结论 ↔ SKU 字段表未回写（评审状态、版本号）。',
     'SKU 台账中评审状态需人工维护，易滞后出错。',
     '评审结论（状态/版本号）自动回写 SKU 字段表，作为后续 SKU 台账的回流字段。',
     '1.5',
     '评审结论变更后 SKU 字段表自动更新；字段口径与 SKU 台账设计一致。'),
]

table = doc.add_table(rows=1, cols=len(COLS))
table.style = 'Table Grid'
table.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = table.rows[0].cells
for i, c in enumerate(COLS):
    hdr[i].text = c
    for p in hdr[i].paragraphs:
        for run in p.runs:
            run.bold = True
            run.font.size = Pt(9)
    set_cell_bg(hdr[i], 'D9E2F3')

PRIO_COLOR = {'P0': 'F4CCCC', 'P1': 'FCE5CD', 'P2': 'FFF2CC', 'D': 'D9EAD3'}

for it in items:
    row = table.add_row().cells
    for i, val in enumerate(it):
        row[i].text = val
        for p in row[i].paragraphs:
            for run in p.runs:
                run.font.size = Pt(8.5)
    set_cell_bg(row[1], PRIO_COLOR[it[1]])
    row[0].paragraphs[0].runs[0].bold = True

for i, w in enumerate(WIDTHS):
    for row in table.rows:
        row.cells[i].width = w

# ---------- 汇总 ----------
doc.add_paragraph()
add_heading('工时与优先级汇总（硬件评审章）', 2)
from collections import OrderedDict
agg = OrderedDict((k, [0, 0.0]) for k in ['P0', 'P1', 'P2', 'D'])
for it in items:
    agg[it[1]][0] += 1
    agg[it[1]][1] += float(it[5])
sum_table = doc.add_table(rows=1, cols=3)
sum_table.style = 'Table Grid'
for i, c in enumerate(['优先级', '条目数', '预估工时合计（人日）']):
    sum_table.rows[0].cells[i].text = c
    for p in sum_table.rows[0].cells[i].paragraphs:
        for run in p.runs:
            run.bold = True
    set_cell_bg(sum_table.rows[0].cells[i], 'D9E2F3')
for k, (n, d) in agg.items():
    cells = sum_table.add_row().cells
    cells[0].text = k
    cells[1].text = str(n)
    cells[2].text = '%.1f' % d
tot = sum_table.add_row().cells
tot[0].text = '合计'
tot[1].text = str(sum(v[0] for v in agg.values()))
tot[2].text = '%.1f' % sum(v[1] for v in agg.values())
for c in tot:
    for p in c.paragraphs:
        for run in p.runs:
            run.bold = True

# ---------- 待补评章节 ----------
doc.add_page_break()
add_heading('第二章  待合并与待补评章节', 1)
add_para('以下内容在合并为整站文档前需补齐：', bold=True)
pending = [
    ('首页评审', '此前已完成评审，格式与本清单一致；需找到前次评审记录后合并入第二章。'),
    ('MES 库位评审', '同上，待合并。'),
    ('合同审批评审', '同上（本清单的量化格式即沿用该次的约定），待合并。'),
    ('硬件评审页截图下半部分', '设计说明以下区域（说明编辑器/抄送人/通知设置等）未在截图中显示，'
     '待用户补充全图后追加评审条目。'),
]
for name, note in pending:
    p = doc.add_paragraph(style='List Bullet')
    r = p.add_run(name + '：')
    r.bold = True
    p.add_run(note)

doc.add_paragraph()
add_para('说明：本章节占位条目不计入工时汇总；合并后需重新编排编号并出 V2.0。')

out = r'D:\Project_env\nec-hub-latest\docs\NEC-HUB_整站评审与量化优化清单_V1.0.docx'
doc.save(out)
print('saved:', out)
